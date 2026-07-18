"""Regenerate the synthetic resume parser fixtures used by unit tests."""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


OUTPUT = Path(__file__).parent


def create_text_pdf() -> None:
    pdf = canvas.Canvas(str(OUTPUT / "synthetic-text-resume.pdf"), pagesize=letter)
    lines = [
        "Synthetic Student",
        "synthetic.student@example.invalid | 202-555-0100",
        "Education",
        "Example University, Bachelor of Science in Biology, expected 2028",
        "Experience",
        "Clinical volunteer, 120 hours",
        "Research assistant, molecular biology laboratory",
        "Skills",
        "Patient communication, data analysis, teamwork",
    ]
    y = 740
    for line in lines:
        pdf.drawString(72, y, line)
        y -= 24
    pdf.showPage()
    pdf.drawString(72, 740, "Leadership and service")
    pdf.drawString(72, 716, "Peer mentor and community health volunteer")
    pdf.drawString(72, 692, "Certifications")
    pdf.drawString(72, 668, "Basic Life Support")
    pdf.save()


def create_image_only_like_pdf() -> None:
    pdf = canvas.Canvas(
        str(OUTPUT / "synthetic-image-only-like.pdf"), pagesize=letter
    )
    pdf.setFillGray(0.85)
    pdf.rect(72, 420, 468, 280, fill=1, stroke=0)
    pdf.setFillGray(0.65)
    for row in range(12):
        pdf.rect(96, 660 - row * 18, 360 - (row % 3) * 50, 8, fill=1, stroke=0)
    pdf.save()


def create_docx() -> None:
    document = Document()
    document.add_heading("Synthetic Student", level=1)
    document.add_paragraph("synthetic.student@example.invalid | 202-555-0100")
    document.add_heading("Education", level=2)
    document.add_paragraph(
        "Example University, Bachelor of Science in Biology, expected 2028"
    )
    document.add_heading("Experience", level=2)
    document.add_paragraph("Clinical volunteer, 120 hours")
    document.add_paragraph("Research assistant, molecular biology laboratory")
    document.save(OUTPUT / "synthetic-resume.docx")


if __name__ == "__main__":
    create_text_pdf()
    create_image_only_like_pdf()
    create_docx()
